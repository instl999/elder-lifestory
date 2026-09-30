"""End-to-end pipeline test with the model calls stubbed.

Proves the plumbing between ingest -> extract -> release -> direct -> draft ->
check -> export without spending tokens or needing credentials. What it does
NOT test is prose quality; that is the recognition test in spec v2 §18, and it
needs human readers.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from lifestory import archive as store
from lifestory import compose, confirm, extract, ingest, render, validators
from lifestory.models import (
    Access,
    CaseMeta,
    ConsentRecord,
    ConsentScope,
    InterviewMode,
    Session,
)

FIXTURE = Path(__file__).parent / "fixtures" / "sample-session-1.txt"


@pytest.fixture
def case(tmp_path):
    """A case with a real transcript ingested and consent recorded."""
    (tmp_path / "cases").mkdir()
    meta = CaseMeta(
        id="pipeline",
        owner_name="Margaret Hale",
        owner_preferred_name="Margaret",
        owner_birth_year=1938,
        operator_name="Operator",
        mode=InterviewMode.STANDARD,
    )
    archive, paths = store.create_case(meta, root=tmp_path)

    for scope in (
        ConsentScope.RECORDING,
        ConsentScope.AI_PROCESSING,
        ConsentScope.THIRD_PARTY_SERVICES,
        ConsentScope.FAMILY_SHARING,
        ConsentScope.PRINT,
    ):
        archive.consents.append(
            ConsentRecord(
                scope=scope,
                granted_by="Margaret Hale",
                granted_by_role="story_owner",
                granted=True,
            )
        )

    session = Session(case_id=meta.id, number=1, interviewer="Operator")
    archive.sessions.append(session)
    archive.utterances.extend(
        ingest.load_transcript_file(
            FIXTURE, session, owner_names=["Margaret"], interviewer_names=["Operator"]
        )
    )
    store.save(archive, paths)
    return archive, paths


# ---------------------------------------------------------------------------
# Windowing
# ---------------------------------------------------------------------------


class TestWindowing:
    def test_windows_overlap(self, case):
        archive, _ = case
        windows = extract._windows(archive.utterances, size=10, overlap=3)
        assert len(windows) > 1
        # Consecutive windows share their overlap.
        first_ids = {u.id for u in windows[0]}
        second_ids = {u.id for u in windows[1]}
        assert first_ids & second_ids

    def test_every_utterance_appears(self, case):
        archive, _ = case
        windows = extract._windows(archive.utterances, size=10, overlap=3)
        covered = {u.id for window in windows for u in window}
        assert covered == {u.id for u in archive.utterances}

    def test_passage_exposes_ids_for_citation(self, case):
        archive, _ = case
        passage = extract._format_passage(archive.utterances[:3])
        for utterance in archive.utterances[:3]:
            assert utterance.id in passage
        assert "[owner]" in passage or "[interviewer]" in passage


# ---------------------------------------------------------------------------
# Full pipeline with stubbed model calls
# ---------------------------------------------------------------------------


def _fake_extraction(window) -> extract.ExtractionResult:
    """What a good extraction of the fixture's strongest scene looks like."""
    owner_turns = [u for u in window if u.role.value == "owner"]
    ids = [u.id for u in owner_turns[:3]] or [window[0].id]
    quote_source = next(
        (u for u in window if "not last a winter" in u.text), None
    )
    quotes = []
    if quote_source is not None:
        quotes.append(
            extract.ExtractedQuote(
                text="You'll not last a winter on those wards",
                source_utterance_id=quote_source.id,
            )
        )
    return extract.ExtractionResult(
        story_cards=[
            extract.ExtractedStoryCard(
                title="The job at Worcester",
                period="early adulthood",
                year=1957,
                approximate=True,
                setup="At nineteen she was offered auxiliary nursing at the Royal Infirmary.",
                conflict="Her father did not want her to take it and said she would not last.",
                choice="She did not answer him. She went.",
                consequence="She stayed eleven years.",
                reflection="It was the closest they ever came to a real falling out.",
                people=["Margaret Hale"],
                places=["Worcester"],
                quotes=quotes,
                themes=["independence"],
                strength=5,
                source_utterance_ids=ids,
                followup_question="What did the first winter on the wards actually feel like?",
            )
        ],
        people=[
            extract.ExtractedPerson(
                name="Edith Mary Pargeter", relationship="mother", source_utterance_ids=ids
            )
        ],
        places=[extract.ExtractedPlace(name="Bewdley", source_utterance_ids=ids)],
        claims=[
            extract.ExtractedClaim(
                text="Her mother was Edith Mary Pargeter, born 1908",
                kind="name",
                importance=5,
                source_utterance_ids=ids,
            ),
            extract.ExtractedClaim(
                text="She was born in March 1938",
                kind="date",
                importance=5,
                source_utterance_ids=ids,
            ),
        ],
    )


class TestPipeline:
    def test_extraction_merges_and_verifies_citations(self, case, monkeypatch):
        archive, paths = case
        monkeypatch.setattr(
            extract, "extract_window", lambda window, **kw: _fake_extraction(window)
        )

        session_id = archive.sessions[0].id
        report = extract.extract_session(archive, session_id, window_size=30, overlap=5)

        assert report.story_cards_added > 0
        assert report.claims_added > 0
        assert not report.dropped_citations

        # Every story card traces to a real utterance.
        for card in archive.story_cards:
            assert card.sources
            for ref in card.sources:
                assert archive.utterance(ref.ref_id) is not None

    def test_extraction_refuses_without_consent(self, tmp_path, monkeypatch):
        (tmp_path / "cases").mkdir()
        archive, _ = store.create_case(
            CaseMeta(id="noconsent", owner_name="X"), root=tmp_path
        )
        archive.sessions.append(Session(case_id="noconsent", number=1))
        with pytest.raises(extract.ConsentError):
            extract.extract_session(archive, archive.sessions[0].id)

    def test_cards_start_unreleased(self, case, monkeypatch):
        """Extraction never decides visibility (§12.2)."""
        archive, _ = case
        monkeypatch.setattr(
            extract, "extract_window", lambda window, **kw: _fake_extraction(window)
        )
        extract.extract_session(archive, archive.sessions[0].id, window_size=30)

        assert archive.story_cards
        assert all(c.access is Access.OPERATOR for c in archive.story_cards)
        assert archive.usable_story_cards() == []

    def test_release_then_draft_and_export(self, case, monkeypatch, tmp_path):
        archive, paths = case
        monkeypatch.setattr(
            extract, "extract_window", lambda window, **kw: _fake_extraction(window)
        )
        extract.extract_session(archive, archive.sessions[0].id, window_size=30)

        # Release, the way the CLI does: card and its sources together.
        for card in archive.story_cards:
            card.access = Access.FAMILY
            for ref in card.sources:
                utterance = archive.utterance(ref.ref_id)
                if utterance:
                    utterance.release_to(Access.FAMILY, note="released by Margaret")

        usable = archive.usable_story_cards()
        assert usable

        direction = compose.NarrativeDirection(
            life_theme="She went anyway",
            proposition="That leaving without answering him was the making of her.",
            point_of_view="third_person",
            tone="plain and warm",
            reading_level="accessible",
            structure="chronological",
            chapters=[
                compose.ChapterPlan(
                    number=1,
                    title="The job at Worcester",
                    covers="Leaving home at nineteen against her father's wishes.",
                    story_card_ids=[usable[0].id],
                )
            ],
            sensitive_handling="Her father's doubt is kept, not softened.",
        )

        quote = archive.quotes[0].text if archive.quotes else "she went"
        chapter = compose.DraftedChapter(
            number=1,
            title="The job at Worcester",
            prose=(
                "She was nineteen when the Royal Infirmary offered her the place.\n\n"
                f"Her father did not want her to take it. “{quote},” he told "
                "her, and she did not answer him. She went.\n\n"
                "She stayed eleven years."
            ),
            provenance=[usable[0].id],
        )

        # Validators must pass on honest material.
        report = validators.run_all(archive, draft_text=chapter.prose, for_export=True)
        assert report.ok(), [f.message for f in report.blocking]

        # Markdown manuscript.
        manuscript = render.render_markdown(archive, direction, [chapter])
        assert "The job at Worcester" in manuscript
        assert "Margaret Hale" in manuscript

        # Provenance sheet is operator-facing and names the card.
        provenance = render.render_provenance(archive, [chapter])
        assert usable[0].id in provenance
        assert "Operator use only" in provenance

        # DOCX actually writes.
        out = paths.exports / "memoir.docx"
        render.render_docx(archive, direction, [chapter], out)
        assert out.exists() and out.stat().st_size > 5000

    def test_draft_refuses_unreleased_material(self, case, monkeypatch):
        archive, _ = case
        monkeypatch.setattr(
            extract, "extract_window", lambda window, **kw: _fake_extraction(window)
        )
        extract.extract_session(archive, archive.sessions[0].id, window_size=30)

        # Nothing released -- drafting must refuse rather than quietly use it.
        direction = compose.NarrativeDirection(
            life_theme="x",
            proposition="y",
            point_of_view="third_person",
            tone="plain",
            reading_level="accessible",
            structure="chronological",
            chapters=[
                compose.ChapterPlan(
                    number=1,
                    title="Chapter",
                    covers="z",
                    story_card_ids=[archive.story_cards[0].id],
                )
            ],
            sensitive_handling="n/a",
        )
        with pytest.raises(compose.DraftingError):
            compose.draft_chapter(archive, direction, direction.chapters[0])

    def test_direction_refuses_with_nothing_released(self, case):
        archive, _ = case
        with pytest.raises(compose.DraftingError):
            compose.propose_direction(archive)

    def test_confirmation_queue_from_real_claims(self, case, monkeypatch):
        archive, _ = case
        monkeypatch.setattr(
            extract, "extract_window", lambda window, **kw: _fake_extraction(window)
        )
        extract.extract_session(archive, archive.sessions[0].id, window_size=30)

        # Release first -- the queue never draws on unreleased material (§12.2).
        for utterance in archive.utterances:
            utterance.release_to(Access.FAMILY, note="Margaret, session 2")

        queue = confirm.build(archive)
        assert queue.items
        assert len(queue.items) <= confirm.MAX_ITEMS

        markdown = confirm.render_markdown(queue)
        assert "Margaret" in markdown
        # The family is told explicitly that the book is not theirs to proofread.
        assert "do not need to read or approve the book" in markdown.lower()


# ---------------------------------------------------------------------------
# Persistence
# ---------------------------------------------------------------------------


class TestPersistence:
    def test_save_bumps_revision_and_keeps_history(self, case):
        archive, paths = case
        before = archive.revision
        archive.meta.notes = "a real change, so the save is not a no-op"
        store.save(archive, paths)
        assert archive.revision == before + 1
        assert list(paths.versions.glob("r*.json"))

    def test_round_trip_preserves_everything(self, case):
        archive, paths = case
        store.save(archive, paths)
        reloaded, _ = store.load("pipeline", root=paths.root)
        assert len(reloaded.utterances) == len(archive.utterances)
        assert reloaded.meta.owner_name == archive.meta.owner_name
        assert all(u.access is Access.OWNER_ONLY for u in reloaded.utterances)

    def test_portable_export_is_plain_json(self, case):
        import json

        archive, paths = case
        out = store.export_portable(archive, paths)
        data = json.loads(out.read_text(encoding="utf-8"))
        assert data["schema_version"]
        assert data["meta"]["owner_name"] == "Margaret Hale"


# ---------------------------------------------------------------------------
# A Chinese book, checked the way a reader would notice it
# ---------------------------------------------------------------------------


class TestChineseBook:
    def _book(self, tmp_path, language="zh-Hans"):
        from docx import Document

        (tmp_path / "cases").mkdir()
        archive, paths = store.create_case(
            CaseMeta(id="wang", owner_name="王秀英", owner_preferred_name="王奶奶",
                     language=language),
            root=tmp_path,
        )
        direction = compose.NarrativeDirection(
            life_theme="她还是去了", proposition="没有回答父亲，是她这辈子最勇敢的一次。",
            point_of_view="third_person", tone="朴实温暖", reading_level="通俗",
            structure="chronological", sensitive_handling="父亲的疑虑照原样保留。",
            chapters=[compose.ChapterPlan(number=1, title="去县城", covers="十九岁离家。")],
        )
        chapter = compose.DraftedChapter(
            number=1, title="去县城",
            prose="她十九岁那年，县医院招护理员。\n\n父亲不让她去。她没有回答，走了。",
            provenance=[],
        )
        out = paths.exports / "book.docx"
        render.render_docx(archive, direction, [chapter], out)
        return Document(str(out)), archive, direction, chapter

    def test_chinese_font_is_set_for_chinese_characters(self, tmp_path):
        from docx.oxml.ns import qn

        document, *_ = self._book(tmp_path)
        rfonts = document.styles["Normal"].element.rPr.rFonts
        assert rfonts.get(qn("w:eastAsia")) == "SimSun"

    def test_traditional_uses_a_traditional_font(self, tmp_path):
        from docx.oxml.ns import qn

        document, *_ = self._book(tmp_path, language="zh-Hant")
        assert document.styles["Normal"].element.rPr.rFonts.get(qn("w:eastAsia")) == "PMingLiU"

    def test_front_matter_and_chapter_labels_are_chinese(self, tmp_path):
        document, *_ = self._book(tmp_path)
        text = "\n".join(p.text for p in document.paragraphs)
        assert "王秀英的一生" in text
        assert "目　录" in text
        assert "第一章" in text
        assert "关于本书" in text
        assert "The life of" not in text and "Contents" not in text

    def test_every_body_paragraph_is_indented_two_characters(self, tmp_path):
        from docx.shared import Pt

        document, *_ = self._book(tmp_path)
        body = [p for p in document.paragraphs if p.text.startswith(("她十九岁", "父亲不让"))]
        assert len(body) == 2
        # Chinese convention: the first paragraph is indented too.
        for paragraph in body:
            indent = paragraph.paragraph_format.first_line_indent
            assert indent is None or indent == Pt(24)
        assert document.styles["Normal"].paragraph_format.first_line_indent == Pt(24)

    def test_page_estimate_counts_characters(self, tmp_path):
        _, _, _, chapter = self._book(tmp_path)
        long_chapter = compose.DraftedChapter(
            number=1, title="x", prose="字" * 6500, provenance=[]
        )
        # The old estimate counted whitespace: 6,500 characters with no spaces
        # was one "word", so ~7 pages whatever the length. At ~650 characters
        # a page it is ten pages of text plus front matter and opener.
        pages = render.estimate_pages([long_chapter], "zh")
        assert 16 <= pages <= 19
        doubled = compose.DraftedChapter(number=1, title="x", prose="字" * 13000, provenance=[])
        assert render.estimate_pages([doubled], "zh") > pages + 8

    def test_markdown_manuscript_is_chinese(self, tmp_path):
        _, archive, direction, chapter = self._book(tmp_path)
        text = render.render_markdown(archive, direction, [chapter])
        assert "## 第一章　去县城" in text
        assert "关于本书" in text

    def test_document_properties_name_the_book_not_the_library(self, tmp_path):
        document, *_ = self._book(tmp_path)
        properties = document.core_properties
        assert properties.title == "她还是去了"
        assert properties.author == "王奶奶"
        assert "python-docx" not in f"{properties.author}{properties.comments}"

    def test_counts_dates_and_colons_are_written_the_chinese_way(self, tmp_path):
        from datetime import date

        from lifestory.i18n import format_date

        _, archive, direction, chapter = self._book(tmp_path)
        archive.sessions += [Session(case_id="wang", number=1), Session(case_id="wang", number=2)]
        text = render.render_markdown(archive, direction, [chapter])
        assert "前后共录音两次" in text  # not 2次, not 二次
        assert format_date(date.today(), "zh") in text
        assert date.today().isoformat() not in text

        page = compose.render_direction(direction, "王奶奶", "zh")
        assert "视角：第三人称" in page
        assert ": " not in page.split("## ")[2]  # the how-it's-told block
